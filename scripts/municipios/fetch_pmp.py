#!/usr/bin/env python3
"""Periodo Medio de Pago a proveedores de las EE.LL. (Ministerio de Hacienda).

Portal WebForms (ASP.NET) que requiere cookies + __VIEWSTATE. Flujo:
  1. GET consulta.aspx?tipoPublicacion=2 (año y vista previa)
  2. POST año -> tabla de ficheros "Descargar"
  3. GET descarga.aspx?ejercicio=&periodo=&fichero=

Se descarga "fichero2" (Cesión y variables, trimestral: el más completo, une
modelo de cesión y modelo de variables) del último periodo disponible. Solo se
conservan filas con Tipo de Entidad = "Ayuntamiento"; el resto (mancomunidades,
diputaciones, EATIM, comarcas...) se descarta y se cuenta.

TRAMPA: el código de entidad 'PP-CC-MMM-...' NO es el código INE (el nº de
municipio va dentro de comarca y colisiona entre comarcas). El cruce con
data/municipios.json se hace por PROVINCIA + NOMBRE normalizados, con
fallback a nombre único nacional para los bilingües ('Elche/Elx').
Produce investigacion_cci08/pmp.json: {codigo_ine: {"pmp": dias}}
"""
import html as htmllib
import http.cookiejar
import json
import re
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "investigacion_cci08"
RAW = OUT / "raw"
BASE = ("https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PMP_NET/"
        "aspx/consulta/consulta.aspx?tipoPublicacion=2")
UA = {"User-Agent": "Mozilla/5.0 (mapa-vivienda-espana)"}
TRIM = {"Marzo": "T1", "Junio": "T2", "Septiembre": "T3", "Diciembre": "T4"}

# Nombre de provincia tal como aparece en el fichero PMP (normalizado) -> código INE
PROVINCIAS = {
    "ALMERIA": "04", "CADIZ": "11", "CORDOBA": "14", "GRANADA": "18",
    "HUELVA": "21", "JAEN": "23", "MALAGA": "29", "SEVILLA": "41",
    "HUESCA": "22", "TERUEL": "44", "ZARAGOZA": "50", "ASTURIAS": "33",
    "BALEARES (ILLES)": "07", "ILLES BALEARS": "07", "CANTABRIA": "39",
    "ALBACETE": "02", "CIUDAD REAL": "13", "CUENCA": "16",
    "GUADALAJARA": "19", "TOLEDO": "45", "AVILA": "05", "BURGOS": "09",
    "LEON": "24", "PALENCIA": "34", "SALAMANCA": "37", "SEGOVIA": "40",
    "SORIA": "42", "VALLADOLID": "47", "ZAMORA": "49", "BARCELONA": "08",
    "GIRONA": "17", "LLEIDA": "25", "TARRAGONA": "43",
    "ALICANTE": "03", "ALICANTE/ALACANT": "03", "CASTELLON": "12",
    "CASTELLON/CASTELLO": "12", "VALENCIA": "46", "VALENCIA/VALENCIA": "46",
    "BADAJOZ": "06", "CACERES": "10", "A CORUÑA": "15", "LUGO": "27",
    "OURENSE": "32", "PONTEVEDRA": "36", "MADRID": "28", "MURCIA": "30",
    "NAVARRA": "31", "ARABA/ALAVA": "01", "ARABA": "01", "BIZKAIA": "48",
    "GIPUZKOA": "20", "RIOJA (LA)": "26", "LA RIOJA": "26", "CEUTA": "51",
    "MELILLA": "52", "LAS PALMAS": "35", "PALMAS, LAS": "35",
    "SANTA CRUZ DE TENERIFE": "38", "S.C.TENERIFE": "38",
}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", str(s).upper())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Z0-9]", "", s)


PROV_NORM = {norm(k): v for k, v in PROVINCIAS.items()}


def opener():
    cj = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    op.addheaders = list(UA.items())
    return op


def fields(s: str) -> dict:
    out = {}
    for f in ["__VIEWSTATE", "__VIEWSTATEGENERATOR", "__EVENTVALIDATION"]:
        m = re.search(r'name="%s"[^>]*value="([^"]*)"' % f, s)
        if m:
            out[f] = htmllib.unescape(m.group(1))
    return out


def latest_file() -> tuple[str, str, str]:
    """(periodo, query del fichero2, año). Fichero2 = 'Cesión y variables'."""
    op = opener()
    s = op.open(BASE).read().decode("utf-8", "replace")
    year = re.search(r'lista_ejercicios"[^>]*>\s*<option selected="selected" value="(\d{4})"', s)
    year = year.group(1) if year else None
    m = re.search(r'name="ctl00\$idControlMaster"[^>]*value="([^"]*)"', s)
    data = fields(s)
    data.update({
        "ctl00$idControlMaster": m.group(1) if m else "",
        "__EVENTTARGET": "", "__EVENTARGUMENT": "",
        "ctl00$MainContentPlaceHolder$lista_ejercicios": year,
        "ctl00$MainContentPlaceHolder$boton_consultar": "Consultar",
    })
    r = op.open(BASE, urllib.parse.urlencode(data).encode()).read().decode("utf-8", "replace")
    links = re.findall(
        r"href\s*=\s*'descarga\.aspx\?ejercicio=(\d+)&periodo=([\wñ]+)"
        r"&tipoPublicacion=2&fichero=fichero2'", r)
    if not links:
        raise RuntimeError("sin ficheros fichero2 publicados")
    order = ["Diciembre", "Septiembre", "Junio", "Marzo"]
    links.sort(key=lambda x: order.index(x[1]) if x[1] in order else 99)
    y, p = links[0]
    q = f"ejercicio={y}&periodo={urllib.parse.quote(p)}&tipoPublicacion=2&fichero=fichero2"
    return p, q, y


def main():
    OUT.mkdir(exist_ok=True)
    RAW.mkdir(exist_ok=True)
    periodo, q, year = latest_file()
    tag = TRIM.get(periodo, periodo)
    print(f"periodo: {periodo} {year}")
    raw = RAW / f"pmp_{year}{tag}_cesion_variables.xlsx"
    if not raw.exists():
        url = ("https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PMP_NET/"
               "aspx/consulta/descarga.aspx?" + q)
        raw.write_bytes(opener().open(url).read())

    # índices de cruce: (prov, nombre) y nombre único nacional
    items = json.load(open(ROOT / "data" / "municipios.json"))["items"]
    by_prov_name = {}
    by_name = {}
    for k, v in items.items():
        by_prov_name[(v["p"], norm(v["n"]))] = k
        by_name.setdefault(norm(v["n"]), []).append(k)
    unique_name = {n: ks[0] for n, ks in by_name.items() if len(ks) == 1}

    wb = openpyxl.load_workbook(raw, read_only=True)
    src, descartados, sin_cruce = {}, {}, []
    for sheet in ["Cesión", "Variables"]:
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(values_only=True):
            if r[4] != "Ayuntamiento":
                if r[4] and r[4] != "Tipo de Entidad" and r[3] and "-" in str(r[3]):
                    descartados[r[4]] = descartados.get(r[4], 0) + 1
                continue
            prov = PROV_NORM.get(norm(r[2] or ""))
            nombre = norm(r[5] or "")
            code = by_prov_name.get((prov, nombre))
            if code is None:
                # nombres bilingües 'X/Y' y sufijos entre paréntesis
                for parte in [p for p in re.split(r"[/()]", nombre) if p]:
                    code = by_prov_name.get((prov, parte)) or unique_name.get(parte)
                    if code:
                        break
            if code is None:
                code = unique_name.get(nombre)
            if code is None:
                sin_cruce.append((r[2], str(r[5]).strip()))
                continue
            try:
                src[code] = round(float(r[10]), 1)
            except (TypeError, ValueError):
                continue

    result = {code: {"pmp": p} for code, p in src.items()}
    (OUT / "pmp.json").write_text(json.dumps(result, ensure_ascii=False))
    (OUT / "pmp_meta.json").write_text(json.dumps({
        "period": f"{tag} {year}" if tag.startswith("T") else f"{periodo} {year}",
        "url": BASE, "raw": raw.name, "unit": "días",
        "matched": len(result), "total_source_aytos": len(result) + len(sin_cruce),
        "sin_cruce": sin_cruce,
        "descartados_no_ayuntamiento": descartados,
    }, ensure_ascii=False, indent=1))
    print(f"pmp.json: {len(result)} municipios; sin cruce: {len(sin_cruce)}")
    for x in sin_cruce[:15]:
        print("  sin cruce:", x)
    print(f"entidades descartadas (no ayuntamiento): {descartados}")


if __name__ == "__main__":
    main()
